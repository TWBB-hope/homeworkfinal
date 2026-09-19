"""生成 docs/期末大作业报告.docx
以教师下发的 final-assignment.docx 为底（保留封面与作业要求），把 docs/期末大作业报告.md
渲染成 Word 正文追加在后面。模板里只有 Normal 样式，所以标题、表格边框、代码块都要手工设置。
"""
import io
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

TEMPLATE = 'final-assignment.docx'
SOURCE = 'docs/期末大作业报告.md'
TARGET = 'docs/期末大作业报告.docx'

BODY_FONT_EA = '宋体'
BODY_FONT_LATIN = 'Times New Roman'
HEAD_FONT = '微软雅黑'
MONO_FONT = 'Consolas'
SCREENSHOT_COLOR = RGBColor(0xB0, 0x5A, 0x00)


def set_fonts(run, ea=BODY_FONT_EA, latin=BODY_FONT_LATIN, size=10.5, bold=False, italic=False, color=None, mono=False):
    if mono:
        ea = latin = MONO_FONT
    run.font.name = latin
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = rpr.makeelement(qn('w:rFonts'), {})
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), ea)
    rfonts.set(qn('w:ascii'), latin)
    rfonts.set(qn('w:hAnsi'), latin)


INLINE = re.compile(r'(\*\*.+?\*\*|`[^`]+`)')
CODE_ONLY = re.compile(r'(`[^`]+`)')


def add_inline(par, text, size=10.5, base_bold=False, mono_code=False):
    """把 **粗体**、`代码`、【截图待补：…】渲染成带格式的 run。
    粗体里还会嵌行内代码（如 **`js/data.js`**），所以粗体分支要再拆一层。"""
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith('**') and piece.endswith('**'):
            for sub in CODE_ONLY.split(piece[2:-2]):
                if not sub:
                    continue
                run = par.add_run(sub[1:-1] if sub.startswith('`') else sub)
                set_fonts(run, size=size, bold=True, mono=sub.startswith('`'))
        elif piece.startswith('`') and piece.endswith('`'):
            run = par.add_run(piece[1:-1])
            set_fonts(run, size=size - 0.5, bold=base_bold, mono=True)
        else:
            run = par.add_run(piece)
            set_fonts(run, size=size, bold=base_bold, mono=mono_code)
        if '【截图待补' in piece:
            run.font.color.rgb = SCREENSHOT_COLOR
            run.bold = True


def add_heading(doc, text, level):
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt({1: 22, 2: 18, 3: 12, 4: 10}[level])
    par.paragraph_format.space_after = Pt({1: 10, 2: 8, 3: 6, 4: 4}[level])
    par.paragraph_format.keep_with_next = True
    size = {1: 16, 2: 14, 3: 12, 4: 11}[level]
    run = par.add_run(text.replace('`', ''))
    set_fonts(run, ea=HEAD_FONT, latin=HEAD_FONT, size=size, bold=True)
    if level >= 3:
        run.font.color.rgb = RGBColor(0x1F, 0x5F, 0x5B)
    return par


def add_body(doc, text):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(4)
    par.paragraph_format.line_spacing = 1.28
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline(par, text)
    return par


def add_bullet(doc, text, order=None):
    par = doc.add_paragraph()
    par.paragraph_format.left_indent = Pt(18)
    par.paragraph_format.space_after = Pt(2)
    par.paragraph_format.line_spacing = 1.24
    prefix = (order + '  ') if order else '•  '
    run = par.add_run(prefix)
    set_fonts(run)
    add_inline(par, text)
    return par


def add_code_block(doc, lines):
    for line in lines:
        par = doc.add_paragraph()
        par.paragraph_format.left_indent = Pt(12)
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.line_spacing = 1.0
        run = par.add_run(line if line else ' ')
        set_fonts(run, size=9, mono=True)
        run.font.color.rgb = RGBColor(0x33, 0x3A, 0x3D)
        ppr = par._element.get_or_add_pPr()
        shd = ppr.makeelement(qn('w:shd'), {qn('w:val'): 'clear', qn('w:fill'): 'F2F4F5'})
        ppr.append(shd)


def add_borders(table):
    tbl_pr = table._element.tblPr
    borders = tbl_pr.makeelement(qn('w:tblBorders'), {})
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = borders.makeelement(qn('w:' + edge), {
            qn('w:val'): 'single', qn('w:sz'): '4', qn('w:color'): 'A9B7B3'})
        borders.append(el)
    tbl_pr.append(borders)


def add_table(doc, rows):
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    add_borders(table)
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for col in range(cols):
            text = row[col] if col < len(row) else ''
            par = cells[col].paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.line_spacing = 1.1
            add_inline(par, text, size=8.5, base_bold=(index == 0))
            if index == 0:
                shd = par._element.get_or_add_pPr().makeelement(
                    qn('w:shd'), {qn('w:val'): 'clear', qn('w:fill'): 'E4EBE8'})
                par._element.get_or_add_pPr().append(shd)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def split_row(line):
    return [c.strip() for c in line.strip().strip('|').split('|')]


def is_sep(line):
    return bool(re.match(r'^\|[\s:|-]+\|$', line.strip()))


def fill_cover(doc):
    """封面只填能确定的两项：姓名与学号。年级／专业／教师／日期留空由本人手写。"""
    cover = doc.paragraphs[11]
    texts = [r.text for r in cover.runs]
    if texts and texts[0] == '学生姓名':
        cover.runs[0].text = '学生姓名  李偲钿'
    if len(texts) > 2 and texts[2] == '学号':
        cover.runs[2].text = '学号  20251060150'


def trim_outline(doc, keep_until=22):
    """删掉模板里"一~七"的要求提纲，正文会按同样的顺序重排。"""
    for par in list(doc.paragraphs)[keep_until + 1:]:
        par._element.getparent().remove(par._element)


def ensure_section(doc):
    """教师模板的 document.xml 里没有 sectPr，python-docx 计算表格宽度时会取不到节而报错。
    补一个 A4 节属性（页边距按课程论文常用值），表格与分页才能正常排版。"""
    import copy

    if doc.element.sectPr_lst:
        return
    ref = Document()
    doc.element.body.append(copy.deepcopy(ref.sections[-1]._sectPr))
    sec = doc.sections[-1]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.left_margin = sec.right_margin = Mm(25)
    sec.top_margin = sec.bottom_margin = Mm(25)


def main():
    doc = Document(TEMPLATE)
    ensure_section(doc)
    fill_cover(doc)
    trim_outline(doc)

    doc.add_page_break()
    add_heading(doc, '期末大作业报告正文', 1)
    add_body(doc, '以下按作业要求的七个部分逐条作答。截图位置以【截图待补：文件名】标出，'
                  '命名与拍摄要求见仓库 screenshots/README.md。')

    lines = io.open(SOURCE, encoding='utf-8').read().split('\n')
    start = next(i for i, l in enumerate(lines) if l.startswith('## 一、'))
    i = start
    table_buf = []
    code_buf = None
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if code_buf is not None:
            if stripped.startswith('```'):
                add_code_block(doc, code_buf)
                code_buf = None
            else:
                code_buf.append(line[2:] if line.startswith('  ') else line)
            i += 1
            continue
        if stripped.startswith('```'):
            code_buf = []
            i += 1
            continue

        if stripped.startswith('|'):
            if is_sep(stripped):
                i += 1
                continue
            table_buf.append(split_row(stripped))
            if i + 1 >= len(lines) or not lines[i + 1].strip().startswith('|'):
                add_table(doc, table_buf)
                table_buf = []
            i += 1
            continue
        if table_buf:
            add_table(doc, table_buf)
            table_buf = []

        if stripped.startswith('#### '):
            add_heading(doc, stripped[5:], 4)
        elif stripped.startswith('### '):
            add_heading(doc, stripped[4:], 3)
        elif stripped.startswith('## '):
            add_heading(doc, stripped[3:], 2)
        elif stripped.startswith('# '):
            add_heading(doc, stripped[2:], 2)
        elif stripped.startswith('---'):
            pass
        elif stripped.startswith('> '):
            par = add_body(doc, stripped[2:])
            par.paragraph_format.left_indent = Pt(18)
            for run in par.runs:
                run.italic = True
        elif re.match(r'^[-*] ', stripped):
            add_bullet(doc, stripped[2:])
        elif re.match(r'^\d+\. ', stripped):
            order, rest = stripped.split('. ', 1)
            add_bullet(doc, rest, order=order + '.')
        elif stripped:
            add_body(doc, stripped)
        i += 1

    if table_buf:
        add_table(doc, table_buf)

    doc.save(TARGET)
    print('已生成', TARGET, '段落数', len(doc.paragraphs), '表格数', len(doc.tables))


if __name__ == '__main__':
    sys.exit(main())
